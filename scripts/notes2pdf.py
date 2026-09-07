#!/usr/bin/env python3
"""강의 필기 JSON(lectures/notes/<id>.json) → 공부용 정리본 PDF (가로, 구간마다 1페이지).
사용: python3 scripts/notes2pdf.py lectures/notes/<id>.json lectures/pdf/<id>.pdf
python-docx로 .docx를 만들고 LibreOffice(soffice)로 PDF 변환 (한글 굵게 지원).
"""
import json, sys, re, os, subprocess, tempfile, shutil
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'Noto Sans CJK KR'
ACCENT = RGBColor(0x37, 0x8A, 0xDD); MUTED = RGBColor(0x6b, 0x72, 0x80); DARK = RGBColor(0x11, 0x18, 0x27); WARN = RGBColor(0x92, 0x40, 0x0e); GREY = RGBColor(0x9c, 0xa3, 0xaf)

def fmt_t(sec):
    sec = int(sec or 0); return f'{sec//60}:{sec%60:02d}'

def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for k in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'): rf.set(qn(k), FONT)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if color is not None: run.font.color.rgb = color

def add_md(par, text, size=12, color=DARK, base_bold=False):
    """**굵게** 마크업을 run으로. 줄바꿈 지원."""
    parts = re.split(r'(\*\*.+?\*\*)', str(text or ''))
    for p in parts:
        if not p: continue
        bold = p.startswith('**') and p.endswith('**')
        chunk = p[2:-2] if bold else p
        lines = chunk.split('\n')
        for i, ln in enumerate(lines):
            if i: par.add_run().add_break()
            r = par.add_run(ln); set_font(r, size, bold or base_bold, color)

def para(doc, text='', size=12, color=DARK, bold=False, space_after=4, md=True, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.space_before = Pt(0)
    if text:
        if md: add_md(p, text, size, color, bold)
        else: r = p.add_run(text); set_font(r, size, bold, color)
    return p

def shade(cell, hex_fill):
    tcPr = cell._element.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_fill); tcPr.append(shd)

def build(note_path, out_path):
    d = json.load(open(note_path, encoding='utf-8'))
    title = d.get('title') or '강의'; course = re.sub(r'^\(SDGs\)', '', d.get('course') or '')
    online = d.get('mode') == 'online'; unit = '구간' if online else '슬라이드'
    slides = d.get('slides', [])

    doc = Document()
    sec = doc.sections[0]; sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.8); sec.top_margin = Cm(1.4); sec.bottom_margin = Cm(1.2)
    # 기본 폰트
    st = doc.styles['Normal']; st.font.name = FONT; st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    # 바닥글
    fp = sec.footer.paragraphs[0]; r = fp.add_run(f'{course} · {title} · 정리본 (Dayble)'); set_font(r, 8.5, False, MUTED)

    # 표지
    para(doc, '', space_after=40)
    para(doc, title, 26, DARK, True, 6, md=False)
    meta = f"{course} · {d.get('date','')} · {unit} {len(slides)}개" + (f" · {fmt_t(d.get('duration'))}" if d.get('duration') else '')
    para(doc, meta, 12, MUTED, False, 16, md=False)
    if d.get('overview'):
        para(doc, '개요', 10, ACCENT, True, 2, md=False); para(doc, d['overview'], 12.5, DARK, False, 12)
    if d.get('exam'):
        para(doc, '시험·과제·출석 언급', 10, ACCENT, True, 2, md=False)
        for x in d['exam']: para(doc, '⚠ ' + x, 12, WARN, False, 2)
        para(doc, '', space_after=8)
    para(doc, '영상을 보면서 이 정리본의 시간표시(0:00)를 따라가면 어디를 듣고 있는지 알 수 있어요. 굵은 글씨가 핵심 용어.', 10.5, MUTED, False, 0, md=False)

    for s in slides:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        segs = s.get('segs') or []
        t0 = fmt_t(segs[0]['t']) if segs else fmt_t(s.get('start', 0)); t1 = fmt_t(segs[-1]['t']) if segs else ''
        h = para(doc, '', space_after=0)
        r = h.add_run(f"{unit} {s.get('page')}"); set_font(r, 16, True, DARK)
        if online or segs:
            r2 = h.add_run(f"   {t0}{'~'+t1 if t1 else ''}"); set_font(r2, 11, False, MUTED)
        if s.get('summary'): para(doc, s['summary'], 11, MUTED, False, 6, md=False)
        else: para(doc, '', space_after=4)
        kps = s.get('keyPoints') or (s.get('liveNotes') or {}).get('notes') or []
        if kps:
            for k in kps:
                p = para(doc, '', 12.5, DARK, False, 3); p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.first_line_indent = Cm(-0.45)
                add_md(p, '•  ' + k, 12.5, DARK)
        else:
            para(doc, '핵심 없음 (설명 없이 넘김)', 11, MUTED, False, 4, md=False)
        # 대본
        rows = []
        for g in segs:
            txt = g.get('fix') or g.get('ko') or g.get('en') or ''
            if not txt: continue
            orig = g.get('fix') or g.get('en') or ''
            rows.append((fmt_t(g.get('t')), txt, orig if (orig and orig != txt and g.get('ko')) else ''))
        if not rows and s.get('ko'): rows = [('', s['ko'], '')]
        if rows:
            para(doc, '', space_after=6); para(doc, '대본', 10, ACCENT, True, 2, md=False)
            tbl = doc.add_table(rows=0, cols=2); tbl.autofit = False
            tblPr = tbl._tbl.tblPr; lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
            for ts, txt, orig in rows:
                cells = tbl.add_row().cells
                cells[0].width = Cm(1.6); cells[1].width = Cm(24.3)
                p0 = cells[0].paragraphs[0]; p0.paragraph_format.space_after = Pt(1); r0 = p0.add_run(ts); set_font(r0, 8.5, False, MUTED)
                p1 = cells[1].paragraphs[0]; p1.paragraph_format.space_after = Pt(1); r1 = p1.add_run(txt); set_font(r1, 9.5, False, RGBColor(0x37, 0x41, 0x51))
                if orig:
                    p2 = cells[1].add_paragraph(); p2.paragraph_format.space_after = Pt(2); r2 = p2.add_run(orig); set_font(r2, 8.5, False, GREY)
            for col, w in zip(tbl.columns, (Cm(1.6), Cm(24.3))): col.width = w
            grid = tbl._tbl.find(qn('w:tblGrid'))
            if grid is not None:
                for gc, w in zip(grid.findall(qn('w:gridCol')), (Cm(1.6), Cm(24.3))): gc.set(qn('w:w'), str(int(w.twips)))
    tmp = tempfile.mkdtemp(); docx_path = os.path.join(tmp, 'note.docx'); doc.save(docx_path)
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmp, docx_path], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    os.makedirs(os.path.dirname(out_path) or '.', exist_ok=True); shutil.move(os.path.join(tmp, 'note.pdf'), out_path); shutil.rmtree(tmp, ignore_errors=True)

if __name__ == '__main__':
    src, out = sys.argv[1], sys.argv[2]
    build(src, out); print('pdf:', out, os.path.getsize(out), 'bytes')
